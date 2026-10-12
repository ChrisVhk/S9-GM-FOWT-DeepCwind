#!/usr/bin/env python3
"""Contrôle du classeur d'architecture du flotteur (séance 0b, point 7) contre le fichier HydroDyn du modèle flottant.

Usage :  python scripts/verifier_architecture.py <classeur rempli.xlsx> [--hydrodyn <fichier HydroDyn>]

Pour chacun des 25 membres du fichier HydroDyn de `modeles/flottant/` (celui que cite `HydroFile` dans son `main.fst`), le script
lit dans le classeur : les deux nœuds du membre (coordonnées lues dans la feuille `Noeuds`), son jeu de propriétés (diamètre et
épaisseur lus dans `Proprietes`), sa discrétisation, et le Cd et le Ca de son groupe (feuille `Coefficients`). Il compare à
HydroDyn et écrit une ligne par membre :

    Membre M5 (Base Column 1) : juste | écart (...) | manquant (...)

« manquant » : une des cellules nécessaires est vide ou porte encore « à remplir » (le membre n'est pas jugé). « écart » : toutes
les cellules sont remplies et au moins une diffère du fichier (le champ fautif est nommé, avec votre valeur et celle du fichier).
« juste » : tout concorde. Code de sortie 0 si les 25 membres sont justes, 1 sinon, 2 si le classeur est illisible.

Numérotation libre : les nœuds et les jeux de propriétés sont référencés par leurs ÉTIQUETTES du classeur (`N3`, `P2`), pas par les
numéros du fichier HydroDyn ; seuls les membres sont repérés par leur étiquette `M<numéro HydroDyn>`.
Tolérances : 5 cm sur une coordonnée (les tableaux du document de définition OC5 sont arrondis : jusqu'à 4,2 cm d'écart avec le fichier
HydroDyn), 0,1 mm sur un diamètre ou une épaisseur, 0,001 sur Cd, Ca et la discrétisation.
Ce que le script NE contrôle PAS : l'exactitude physique du fichier HydroDyn lui-même (c'est la référence de l'exercice), ni les
feuilles `Donnees` (version étudiante) et `Sources` (version complète des coulisses)."""
import argparse
import re
import sys
from pathlib import Path

RACINE = Path(__file__).resolve().parents[1]
TOL_COORD, TOL_SECTION, TOL_COEF = 0.05, 1e-4, 1e-3
A_REMPLIR = "à remplir"


# ---------------------------------------------------------------------------------------------------- HydroDyn
def fichier_hydrodyn_du_modele(modele=RACINE / "modeles" / "flottant"):
    for ligne in (modele / "main.fst").read_text(encoding="utf-8").splitlines():
        m = re.match(r'^\s*"([^"]+)"\s+HydroFile\b', ligne)
        if m:
            return (modele / m.group(1)).resolve()
    raise FileNotFoundError(f"HydroFile introuvable dans {modele / 'main.fst'}")


def _table(lignes, titre, n_min):
    i = next((k for k, l in enumerate(lignes) if titre in l), None)
    if i is None:
        raise ValueError(f"section « {titre} » absente du fichier HydroDyn")
    i += 4                                                      # titre, compte, en-tête, unités -> première donnée
    sortie = []
    while i < len(lignes):
        corps, _, commentaire = lignes[i].partition("!")
        cols = corps.split()
        if len(cols) < n_min or not re.match(r"^-?\d", cols[0]):
            break
        sortie.append((cols, commentaire.strip()))
        i += 1
    return sortie


def lire_hydrodyn(chemin):
    """-> {numéro de membre : dict(nom, p1, p2 (coordonnées des joints), D, t, div, cd, ca)}"""
    lignes = Path(chemin).read_text(encoding="utf-8").splitlines()
    joints = {int(c[0]): tuple(float(v) for v in c[1:4]) for c, _ in _table(lignes, "MEMBER JOINTS", 4)}
    jeux = {int(c[0]): (float(c[1]), float(c[2])) for c, _ in _table(lignes, "CYLINDRICAL MEMBER CROSS-SECTION PROPERTIES", 3)}
    coefs = {int(c[0]): (float(c[1]), float(c[5]))
             for c, _ in _table(lignes, "MEMBER-BASED CYLINDRICAL-MEMBER HYDRODYNAMIC COEFFICIENTS (model 3)", 6)}
    membres = {}
    for c, nom in _table(lignes, "-------------------- MEMBERS", 8):
        i, j1, j2, p = int(c[0]), int(c[1]), int(c[2]), int(c[3])
        membres[i] = dict(nom=nom or f"membre {i}", p1=joints[j1], p2=joints[j2], D=jeux[p][0], t=jeux[p][1],
                          div=float(c[7]), cd=coefs[i][0], ca=coefs[i][1])
    return membres


# ---------------------------------------------------------------------------------------------------- classeur
def _nombre(v):
    """Une cellule est un nombre si elle contient un nombre (ou un texte numérique) ; « à remplir », vide : None."""
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def _etiquette(v):
    if v is None:
        return None
    s = str(v).strip()
    return None if not s or s.lower() == A_REMPLIR else s


def lire_classeur(chemin):
    import openpyxl
    wb = openpyxl.load_workbook(chemin, data_only=True)
    for feuille in ("Noeuds", "Membres", "Proprietes", "Coefficients"):
        if feuille not in wb.sheetnames:
            raise ValueError(f"feuille « {feuille} » absente du classeur")
    noeuds = {}
    for r in wb["Noeuds"].iter_rows(min_row=2, values_only=True):
        e = _etiquette(r[0])
        if e:
            noeuds[e] = tuple(_nombre(v) for v in r[1:4])
    jeux = {}
    for r in wb["Proprietes"].iter_rows(min_row=2, values_only=True):
        e = _etiquette(r[0])
        if e:
            jeux[e] = (_nombre(r[1]), _nombre(r[2]))
    coefs = {}
    for r in wb["Coefficients"].iter_rows(min_row=2, values_only=True):
        e = _etiquette(r[0])
        if e:
            coefs[e] = (_nombre(r[1]), _nombre(r[2]))
    membres = {}
    for r in wb["Membres"].iter_rows(min_row=2, values_only=True):
        m = re.match(r"^M(\d+)$", _etiquette(r[0]) or "")
        if m:
            membres[int(m[1])] = dict(groupe=_etiquette(r[2]), n1=_etiquette(r[3]), n2=_etiquette(r[4]),
                                      jeu=_etiquette(r[5]), div=_nombre(r[6]))
    return dict(noeuds=noeuds, jeux=jeux, coefs=coefs, membres=membres)


# ---------------------------------------------------------------------------------------------------- comparaison
def _ecart(a, b, tol):
    return abs(a - b) > tol


def comparer(classeur, hydrodyn):
    """-> liste de (numéro, nom, statut, détail) pour chacun des membres de HydroDyn, dans l'ordre du numéro."""
    resultat = []
    for i in sorted(hydrodyn):
        h = hydrodyn[i]
        w = classeur["membres"].get(i)
        manquants, ecarts = [], []
        if w is None:
            resultat.append((i, h["nom"], "manquant", "ligne absente du classeur"))
            continue
        for k, (cle, ref) in enumerate((("n1", h["p1"]), ("n2", h["p2"])), start=1):
            etiq = w[cle]
            xyz = classeur["noeuds"].get(etiq) if etiq else None
            if etiq is None:
                manquants.append(f"nœud {k}")
            elif xyz is None or any(v is None for v in xyz):
                manquants.append(f"coordonnées de {etiq} (nœud {k})")
            else:
                for axe, v, r in zip("xyz", xyz, ref):
                    if _ecart(v, r, TOL_COORD):
                        ecarts.append(f"nœud {k} {etiq} : {axe} = {v:g} (fichier : {r:g})")
        jeu = w["jeu"]
        sect = classeur["jeux"].get(jeu) if jeu else None
        if jeu is None:
            manquants.append("jeu de propriétés")
        elif sect is None or any(v is None for v in sect):
            manquants.append(f"diamètre ou épaisseur du jeu {jeu}")
        else:
            if _ecart(sect[0], h["D"], TOL_SECTION):
                ecarts.append(f"jeu {jeu} : diamètre = {sect[0]:g} (fichier : {h['D']:g})")
            if _ecart(sect[1], h["t"], TOL_SECTION):
                ecarts.append(f"jeu {jeu} : épaisseur = {sect[1]:g} (fichier : {h['t']:g})")
        if w["div"] is None:
            manquants.append("discrétisation")
        elif _ecart(w["div"], h["div"], TOL_COEF):
            ecarts.append(f"discrétisation = {w['div']:g} (fichier : {h['div']:g})")
        g = w["groupe"]
        cc = classeur["coefs"].get(g) if g else None
        if g is None:
            manquants.append("groupe")
        elif cc is None or cc[0] is None:
            manquants.append(f"Cd du groupe {g}")
        elif _ecart(cc[0], h["cd"], TOL_COEF):
            ecarts.append(f"Cd du groupe {g} = {cc[0]:g} (fichier : {h['cd']:g})")
        if g is not None and cc is not None:
            if cc[1] is None:
                manquants.append(f"Ca du groupe {g}")
            elif _ecart(cc[1], h["ca"], TOL_COEF):
                ecarts.append(f"Ca du groupe {g} = {cc[1]:g} (fichier : {h['ca']:g})")
        if manquants:
            resultat.append((i, h["nom"], "manquant", ", ".join(manquants)))
        elif ecarts:
            resultat.append((i, h["nom"], "écart", " ; ".join(ecarts)))
        else:
            resultat.append((i, h["nom"], "juste", ""))
    return resultat


def main(argv=None):
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("classeur")
    p.add_argument("--hydrodyn", default=None, help="fichier HydroDyn de référence (défaut : celui de modeles/flottant)")
    a = p.parse_args(argv)
    try:
        hyd = lire_hydrodyn(a.hydrodyn or fichier_hydrodyn_du_modele())
    except (OSError, ValueError, KeyError, StopIteration) as e:
        print(f"Fichier HydroDyn illisible : {e}", file=sys.stderr)
        return 2
    try:
        cl = lire_classeur(a.classeur)
    except Exception as e:                       # classeur absent, corrompu ou sans les feuilles attendues
        print(f"Classeur illisible : {a.classeur} ({type(e).__name__} : {e})", file=sys.stderr)
        return 2
    res = comparer(cl, hyd)
    for i, nom, statut, detail in res:
        print(f"Membre M{i} ({nom}) : {statut}" + (f" ({detail})" if detail else ""))
    n = {s: sum(1 for r in res if r[2] == s) for s in ("juste", "écart", "manquant")}
    print(f"\n{len(res)} membres : {n['juste']} juste(s), {n['écart']} écart(s), {n['manquant']} manquant(s)")
    return 0 if n["juste"] == len(res) else 1


if __name__ == "__main__":
    sys.exit(main())
